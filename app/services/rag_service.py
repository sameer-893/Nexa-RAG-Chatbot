import uuid
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any

from app.services.routing_service import routing_service

from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_core.messages import (
    AIMessage,
    HumanMessage,
    SystemMessage,
)
from langchain_openai import ChatOpenAI
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

from app.config import settings
from app.services.pinecone_service import pinecone_service
from app.services.errors import NotFoundError
from app.services.redis_service import make_chat_title, redis_service


# ============================================================
# DEFAULT SYSTEM PROMPT
# ============================================================

DEFAULT_SYSTEM_PROMPT = """
You are a helpful AI assistant.

Answer the user's question accurately and clearly.

The system may provide:
- conversation history,
- retrieved document context,
- or information obtained from web search.

Use the information provided by the selected route.

Rules:

1. Do not invent facts.
2. For RAG questions, prefer the retrieved document context.
3. For web-search questions, rely on the retrieved web information.
4. For direct questions, answer using your general knowledge and reasoning.
5. Use conversation history when it helps understand the user's question.
6. If the required information is unavailable, say so clearly.
7. Keep answers relevant and easy to understand.
8. Mention sources when they are available.
""".strip()


class RAGService:

    def __init__(self):

        # =====================================================
        # TEXT SPLITTER
        # =====================================================

        self.splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.chunk_size,
            chunk_overlap=settings.chunk_overlap,
            separators=[
                "\n\n",
                "\n",
                ". ",
                " ",
                "",
            ],
        )

        # =====================================================
        # NORMAL LLM
        # Used for RAG answers.
        # =====================================================

        self.llm = ChatOpenAI(
            model=settings.chat_model,
            api_key=settings.openai_api_key,
        )

        # =====================================================
        # INITIAL REDIS CONFIG
        # =====================================================

        if redis_service.get_prompt() is None:

            redis_service.set_prompt(
                DEFAULT_SYSTEM_PROMPT
            )

        if redis_service.get_context_limit() is None:

            redis_service.set_context_limit(
                settings.chat_context_messages
            )

    # =========================================================
    # TEXT EXTRACTION
    # =========================================================

    def extract_text(
        self,
        filename: str,
        content: bytes,
    ) -> str:

        extension = (
            Path(filename)
            .suffix
            .lower()
        )

        # -----------------------------------------------------
        # PDF
        # -----------------------------------------------------

        if extension == ".pdf":

            reader = PdfReader(
                BytesIO(content)
            )

            pages: list[str] = []

            for page in reader.pages:

                page_text = page.extract_text()

                if page_text:

                    pages.append(
                        page_text
                    )

            return "\n\n".join(pages)

        # -----------------------------------------------------
        # DOCX
        # -----------------------------------------------------

        if extension == ".docx":

            document = DocxDocument(
                BytesIO(content)
            )

            paragraphs = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            return "\n".join(paragraphs)

        # -----------------------------------------------------
        # TXT / Markdown
        # -----------------------------------------------------

        if extension in {".txt", ".md"}:

            return content.decode(
                "utf-8",
                errors="ignore",
            )

        raise ValueError(
            "Unsupported file type. "
            "Supported: PDF, DOCX, TXT, MD."
        )

    # =========================================================
    # ADD FILE
    # =========================================================

    def add_file(
        self,
        user_id: str,
        filename: str,
        content: bytes,
    ) -> dict[str, Any]:

        text = self.extract_text(
            filename,
            content,
        )

        if not text.strip():

            raise ValueError(
                "No readable text was found in the file."
            )

        file_id = str(
            uuid.uuid4()
        )

        chunks = self.splitter.split_text(
            text
        )

        if not chunks:

            raise ValueError(
                "The file did not produce any text chunks."
            )

        documents: list[Document] = []
        vector_ids: list[str] = []

        for index, chunk in enumerate(
            chunks
        ):

            vector_id = (
                f"{file_id}:{index}"
            )

            metadata = {

                "file_id": file_id,

                "user_id": user_id,

                "filename": filename,

                "chunk_index": index,
            }

            documents.append(
                Document(
                    page_content=chunk,
                    metadata=metadata,
                )
            )

            vector_ids.append(
                vector_id
            )

        # -----------------------------------------------------
        # Store vectors in Pinecone
        # -----------------------------------------------------

        pinecone_service.add_documents(
            documents=documents,
            ids=vector_ids,
        )

        # -----------------------------------------------------
        # Store document metadata in Redis
        # -----------------------------------------------------

        document_record = {

            "file_id": file_id,

            "filename": filename,

            "user_id": user_id,

            "chunk_count": len(chunks),

            "size_bytes": len(content),

            "uploaded_at": datetime.now(
                timezone.utc
            ).isoformat(),

            "sample_vector_id": vector_ids[0],
        }

        redis_service.save_document(
            user_id=user_id,
            file_id=file_id,
            document=document_record,
        )

        return document_record

    # =========================================================
    # REMOVE FILE
    # =========================================================

    def remove_file(
        self,
        user_id: str,
        file_id: str,
    ) -> dict[str, str]:

        document = redis_service.get_document(
            user_id,
            file_id,
        )

        if document is None:

            raise NotFoundError(
                "Document not found."
            )

        # -----------------------------------------------------
        # Delete vectors from Pinecone
        # -----------------------------------------------------

        vector_ids = [
            f"{file_id}:{index}"
            for index in range(
                int(document.get("chunk_count") or 0)
            )
        ]

        pinecone_service.delete_file(
            user_id=user_id,
            file_id=file_id,
            vector_ids=vector_ids,
        )

        # -----------------------------------------------------
        # Delete metadata from Redis
        # -----------------------------------------------------

        redis_service.delete_document(
            user_id=user_id,
            file_id=file_id,
        )

        return {

            "message": (
                "Document removed successfully."
            ),

            "file_id": file_id,
        }

    # =========================================================
    # GET DOCUMENT DETAIL
    # =========================================================

    def get_document(
        self,
        user_id: str,
        file_id: str,
    ) -> dict[str, Any]:

        document = redis_service.get_document(
            user_id,
            file_id,
        )

        if document is None:

            raise NotFoundError(
                "Document not found."
            )

        sample_vector_id = (
            document["sample_vector_id"]
        )

        pinecone_result = (
            pinecone_service.fetch_vectors(
                [sample_vector_id]
            )
        )

        vectors = getattr(
            pinecone_result,
            "vectors",
            {},
        )

        return {

            **document,

            "vector_exists": (
                sample_vector_id in vectors
            ),
        }

    # =========================================================
    # GET ALL DOCUMENTS
    # =========================================================

    def get_documents(
        self,
        user_id: str,
    ) -> list[dict[str, Any]]:

        documents = redis_service.get_documents(
            user_id
        )

        return sorted(
            documents,
            key=lambda item: (
                item.get("uploaded_at") or "",
                item.get("filename") or "",
            ),
        )

    # =========================================================
    # RETRIEVE FROM PINECONE
    # =========================================================

    def retrieve(
        self,
        user_id: str,
        question: str,
    ):

        return pinecone_service.search(
            query=question,
            user_id=user_id,
            top_k=settings.retrieval_top_k,
        )

    # =========================================================
    # BUILD RAG CONTEXT
    # =========================================================

    def build_context(
        self,
        search_results,
    ) -> str:

        if not search_results:

            return (
                "No relevant documents were retrieved."
            )

        parts: list[str] = []

        for position, (
            document,
            score,
        ) in enumerate(
            search_results,
            start=1,
        ):

            filename = document.metadata.get(
                "filename",
                "Unknown",
            )

            parts.append(
                f"""
SOURCE {position}
Filename: {filename}
Similarity score: {score:.4f}

{document.page_content}
""".strip()
            )

        return "\n\n---\n\n".join(
            parts
        )

    # =========================================================
    # CHAT
    # =========================================================

    def chat(
        self,
        user_id: str,
        question: str,
        chat_id: str | None = None,
    ) -> dict[str, Any]:

        question = question.strip()

        if not question:

            raise ValueError(
                "Question cannot be empty."
            )

        # =====================================================
        # 1. GET LLM CONTEXT LIMIT
        #
        # This controls how much history the LLM/router sees.
        # It DOES NOT delete older messages from Redis.
        # =====================================================

        context_limit = (
            redis_service.get_context_limit()
        )

        if context_limit is None:

            context_limit = (
                settings.chat_context_messages
            )

        # A missing chat_id means "start a new chat": it is created
        # only after the answer exists, so a failed request leaves
        # no empty chat behind.
        if chat_id is not None:

            redis_service.get_chat(
                user_id,
                chat_id,
            )  # raises NotFoundError

            recent_history = (
                redis_service.get_recent_history(
                    user_id,
                    chat_id,
                    context_limit,
                )
            )

        else:

            recent_history = []

        # =====================================================
        # 2. GET USER DOCUMENTS
        #
        # Router needs to know whether the user has documents.
        # =====================================================

        user_documents = (
            redis_service.get_documents(
                user_id
            )
        )

        document_names = [
            document.get(
                "filename",
                "Unknown",
            )
            for document in user_documents
        ]

        # =====================================================
        # 3. GET SYSTEM PROMPT
        # =====================================================

        system_prompt = (
            redis_service.get_prompt()
            or DEFAULT_SYSTEM_PROMPT
        )

        # =====================================================
        # 4. AGENTIC ROUTING
        #
        # Possible routes:
        #
        # direct -> normal LLM
        # rag    -> Pinecone + LLM
        # web    -> web search + LLM
        # =====================================================

        decision = (
            routing_service.decide_route(
                question=question,
                history=recent_history,
                document_names=document_names,
            )
        )

        route = decision.route

        # All sources returned to the UI/API
        sources: list[dict[str, Any]] = []

        # =====================================================
        # 5. DIRECT ROUTE
        # =====================================================

        if route == "direct":

            answer = (
                routing_service.direct_answer(
                    question=question,
                    history=recent_history,
                    system_prompt=system_prompt,
                )
            )

        # =====================================================
        # 6. RAG ROUTE
        # =====================================================

        elif route == "rag":

            search_results = self.retrieve(
                user_id=user_id,
                question=question,
            )

            rag_context = (
                self.build_context(
                    search_results
                )
            )

            messages = [

                SystemMessage(
                    content=(
                        f"{system_prompt}\n\n"

                        "The selected route is RAG.\n\n"

                        "Use the following retrieved "
                        "document context to answer "
                        "the user's question.\n\n"

                        f"{rag_context}"
                    )
                )
            ]

            # -------------------------------------------------
            # Add bounded chat history
            # -------------------------------------------------

            for item in recent_history:

                role = item.get(
                    "role"
                )

                content = item.get(
                    "content",
                    "",
                )

                if not content:
                    continue

                if role == "user":

                    messages.append(
                        HumanMessage(
                            content=content
                        )
                    )

                elif role == "assistant":

                    messages.append(
                        AIMessage(
                            content=content
                        )
                    )

            # -------------------------------------------------
            # Current question
            # -------------------------------------------------

            messages.append(
                HumanMessage(
                    content=question
                )
            )

            # -------------------------------------------------
            # Generate RAG answer
            # -------------------------------------------------

            response = self.llm.invoke(
                messages
            )

            answer = self._extract_text(
                response
            )

            # -------------------------------------------------
            # RAG sources
            # -------------------------------------------------

            for (
                document,
                score,
            ) in search_results:

                sources.append(
                    {

                        "type": "document",

                        "filename": (
                            document.metadata.get(
                                "filename"
                            )
                        ),

                        "file_id": (
                            document.metadata.get(
                                "file_id"
                            )
                        ),

                        "chunk_index": (
                            document.metadata.get(
                                "chunk_index"
                            )
                        ),

                        "score": float(score),
                    }
                )

        # =====================================================
        # 7. WEB SEARCH ROUTE
        # =====================================================

        elif route == "web":

            web_result = (
                routing_service.web_answer(
                    question=question,
                    history=recent_history,
                    system_prompt=system_prompt,
                )
            )

            answer = web_result[
                "answer"
            ]

            for source in web_result.get(
                "sources",
                [],
            ):

                sources.append(
                    {
                        "type": "web",
                        **source,
                    }
                )

        # =====================================================
        # 8. INVALID ROUTE
        # =====================================================

        else:

            raise RuntimeError(
                f"Unknown routing decision: {route}"
            )

        # =====================================================
        # 9. SAVE FULL CHAT HISTORY
        #
        # Both messages are written in ONE atomic call. The answer
        # keeps its route and sources so reopening the chat shows
        # them again; only role/content are ever sent to the LLM.
        # =====================================================

        if chat_id is None:

            chat_id = redis_service.create_chat(
                user_id
            )["chat_id"]

        saved_user, saved_assistant = (
            redis_service.add_messages(
                user_id=user_id,
                chat_id=chat_id,
                messages=[
                    {
                        "role": "user",
                        "content": question,
                    },
                    {
                        "role": "assistant",
                        "content": answer,
                        "route": route,
                        "route_reason": decision.reason,
                        "sources": sources,
                    },
                ],
                title_hint=make_chat_title(
                    question
                ),
            )
        )

        # =====================================================
        # 10. COMPLETE HISTORY COUNT
        # =====================================================

        total_messages = (
            redis_service.count_messages(
                user_id,
                chat_id,
            )
        )

        chat_summary = redis_service.get_chat(
            user_id,
            chat_id,
        )

        # =====================================================
        # 11. RETURN RESULT
        # =====================================================

        return {

            "answer": answer,

            "route": route,

            "route_reason": decision.reason,

            "sources": sources,

            "chat_id": chat_id,

            "chat_title": chat_summary["title"],

            "user_message": saved_user,

            "assistant_message": saved_assistant,

            "context_messages_used": len(
                recent_history
            ),

            "history_messages_total": total_messages,
        }

    # =========================================================
    # RESPONSE TEXT HELPER
    # =========================================================

    @staticmethod
    def _extract_text(
        response,
    ) -> str:

        content = response.content

        if isinstance(
            content,
            str,
        ):

            return content.strip()

        if isinstance(
            content,
            list,
        ):

            parts: list[str] = []

            for block in content:

                if isinstance(
                    block,
                    str,
                ):

                    parts.append(
                        block
                    )

                elif isinstance(
                    block,
                    dict,
                ):

                    if block.get(
                        "type"
                    ) == "text":

                        text = block.get(
                            "text",
                            "",
                        )

                        if text:

                            parts.append(
                                text
                            )

            return "\n".join(
                parts
            ).strip()

        return str(
            content
        )

    # =========================================================
    # CHATS
    # =========================================================

    def create_chat(
        self,
        user_id: str,
        title: str | None = None,
    ) -> dict[str, Any]:

        return redis_service.create_chat(
            user_id,
            title,
        )

    def list_chats(
        self,
        user_id: str,
    ) -> list[dict[str, Any]]:

        return redis_service.list_chats(
            user_id
        )

    def search_chats(
        self,
        user_id: str,
        query: str,
    ) -> list[dict[str, Any]]:

        return redis_service.search_chats(
            user_id,
            query,
        )

    def get_chat(
        self,
        user_id: str,
        chat_id: str,
    ) -> dict[str, Any]:
        """Chat metadata plus its COMPLETE message history."""

        chat = redis_service.get_chat(
            user_id,
            chat_id,
        )

        return {
            "chat": chat,
            "messages": redis_service.get_messages(
                user_id,
                chat_id,
            ),
        }

    def clear_chat(
        self,
        user_id: str,
        chat_id: str,
    ) -> None:

        redis_service.clear_chat(
            user_id,
            chat_id,
        )

    def delete_chat(
        self,
        user_id: str,
        chat_id: str,
    ) -> None:

        redis_service.delete_chat(
            user_id,
            chat_id,
        )

    def delete_all_chats(
        self,
        user_id: str,
    ) -> int:

        return redis_service.delete_all_chats(
            user_id
        )

    # =========================================================
    # RESPONSES (edit / restore)
    # =========================================================

    def update_response(
        self,
        user_id: str,
        chat_id: str,
        message_id: str,
        content: str,
    ) -> dict[str, Any]:

        content = content.strip()

        if not content:

            raise ValueError(
                "Response cannot be empty."
            )

        return redis_service.update_response(
            user_id,
            chat_id,
            message_id,
            content,
        )

    def restore_response(
        self,
        user_id: str,
        chat_id: str,
        message_id: str,
    ) -> dict[str, Any]:

        return redis_service.restore_response(
            user_id,
            chat_id,
            message_id,
        )

    # =========================================================
    # PROMPT
    # =========================================================

    def get_prompt(
        self,
    ) -> str:

        return (
            redis_service.get_prompt()
            or DEFAULT_SYSTEM_PROMPT
        )

    def update_prompt(
        self,
        prompt: str,
    ) -> str:

        prompt = prompt.strip()

        if not prompt:

            raise ValueError(
                "Prompt cannot be empty."
            )

        redis_service.set_prompt(
            prompt
        )

        return prompt

    def reset_prompt(
        self,
    ) -> str:

        redis_service.set_prompt(
            DEFAULT_SYSTEM_PROMPT
        )

        return DEFAULT_SYSTEM_PROMPT

    # =========================================================
    # CONTEXT LIMIT
    # =========================================================

    def get_context_limit(
        self,
    ) -> int:

        value = (
            redis_service.get_context_limit()
        )

        return (
            value
            if value is not None
            else settings.chat_context_messages
        )

    def update_context_limit(
        self,
        value: int,
    ) -> int:

        if value < 0:

            raise ValueError(
                "Context limit cannot be negative."
            )

        redis_service.set_context_limit(
            value
        )

        return value


# ============================================================
# SINGLE SERVICE INSTANCE
# ============================================================

rag_service = RAGService()